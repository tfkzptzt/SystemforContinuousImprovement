# -*- coding: utf-8 -*-
"""批量导入：Excel(openpyxl read_only=True 流式) / CSV(标准库)，行级校验，逐行成功/失败清单"""
import csv
import io
import os

import config
from db import now

COLUMNS = config.IMPORT_COLUMNS


class ImportError_(Exception):
    """导入文件级别错误（如表头不符）"""


# ---------- 模板 ----------

def ensure_import_template():
    """确保导入模板存在（幂等）"""
    os.makedirs(config.TEMPLATE_DIR, exist_ok=True)
    path = os.path.join(config.TEMPLATE_DIR, config.IMPORT_TEMPLATE_FILENAME)
    if os.path.exists(path):
        return path
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '导入模板'
    ws.append(COLUMNS)
    ws.append(['teacher01', 'GCTZ1001', '2024-2025学年', '第一学期',
               '工程制图课程持续改进报告', '加强实践教学环节，更新实验案例',
               '下一学年通过实验报告评分与学生达成度指标验证', '基本达成目标'])
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[chr(64 + i)].width = 22
    wb.save(path)
    return path


# ---------- 解析 ----------

def _cell(v):
    return '' if v is None else str(v).strip()


def _map_header(header_row):
    header = [_cell(c) for c in header_row]
    idx = {}
    for i, name in enumerate(header):
        if name in COLUMNS and name not in idx:
            idx[name] = i
    missing = [c for c in COLUMNS if c not in idx]
    if missing:
        raise ImportError_('模板列缺失：' + '、'.join(missing))
    return idx


def parse_xlsx(stream):
    """read_only 流式解析，返回 (idx, data_rows)，data_rows: [(行号, [值...])]"""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception as e:
        raise ImportError_(f'Excel 文件解析失败：{e}') from e
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ImportError_('Excel 文件为空') from None
        idx = _map_header(header_row)
        data = []
        row_no = 1
        for values in rows:
            row_no += 1
            cells = [_cell(values[i]) if i < len(values) else '' for i in idx.values()]
            if not any(cells):
                continue
            data.append((row_no, cells))
        return idx, data
    finally:
        wb.close()


def parse_csv(stream):
    raw = stream.read()
    text = raw.decode('utf-8-sig', errors='replace')
    reader = csv.reader(io.StringIO(text))
    all_rows = [r for r in reader]
    if not all_rows:
        raise ImportError_('CSV 文件为空')
    idx = _map_header(all_rows[0])
    data = []
    for row_no, values in enumerate(all_rows[1:], start=2):
        cells = [values[i].strip() if i < len(values) else '' for i in idx.values()]
        if not any(cells):
            continue
        data.append((row_no, cells))
    return idx, data


# ---------- 导入 ----------

def import_rows(db, stream, filename):
    """批量导入，写入 reports(source='import', status='concluded', 含整体结论) +
    measures + 按措施一对一的 achievements(status='approved')。

    :return: {'success_rows': int, 'failed_rows': int, 'errors': [{'row': 行号, 'reason': 原因}]}
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        _, data = parse_xlsx(stream)
    elif ext == '.csv':
        _, data = parse_csv(stream)
    else:
        raise ImportError_('仅支持 .xlsx / .csv 文件')

    result = {'success_rows': 0, 'failed_rows': 0, 'errors': []}

    teachers = {r['username']: r for r in
                db.execute('SELECT id, username FROM users '
                           'WHERE is_teacher=1 AND is_disabled=0').fetchall()}
    courses = {(r['code'], r['academic_year'], r['term']): r['id'] for r in
                db.execute('SELECT id, code, academic_year, term FROM courses').fetchall()}

    # 查重：(教师用户名, 课程代码, 学年, 学期) 已有报告（含历史导入/上传）或本批次内重复均计失败
    existing = set()
    for r in db.execute(
            'SELECT u.username, c.code, c.academic_year, c.term FROM reports r '
            'JOIN users u ON r.teacher_id=u.id JOIN courses c ON r.course_id=c.id').fetchall():
        existing.add((r['username'], r['code'], r['academic_year'], r['term']))
    seen_in_batch = set()

    ts = now()
    for row_no, cells in data:
        rec = dict(zip(COLUMNS, cells))
        reason = None
        if not rec['教师工号']:
            reason = '教师工号为空'
        elif rec['教师工号'] not in teachers:
            reason = f"教师工号「{rec['教师工号']}」不存在"
        elif not rec['课程代码']:
            reason = '课程代码为空'
        elif (rec['课程代码'], rec['学年'], rec['学期']) not in courses:
            reason = f"课程不存在（代码 {rec['课程代码']} / {rec['学年']} / {rec['学期']}）"
        elif (rec['教师工号'], rec['课程代码'], rec['学年'], rec['学期']) in existing \
                or (rec['教师工号'], rec['课程代码'], rec['学年'], rec['学期']) in seen_in_batch:
            reason = '重复记录'
        elif not rec['措施内容']:
            reason = '措施内容为空'
        elif rec['结论'] not in config.CONCLUSION_OPTIONS:
            reason = f"结论「{rec['结论']}」不在可选范围（{'/'.join(config.CONCLUSION_OPTIONS)}）"

        if reason:
            result['failed_rows'] += 1
            result['errors'].append({'row': row_no, 'reason': reason})
            continue

        teacher = teachers[rec['教师工号']]
        course_id = courses[(rec['课程代码'], rec['学年'], rec['学期'])]
        title = rec['报告标题'] or f"{rec['课程代码']}持续改进报告（导入）"

        cur = db.execute(
            'INSERT INTO reports (teacher_id, course_id, title, file_path, content_text, '
            'status, source, conclusion, concluded_at, created_at, updated_at) '
            'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (teacher['id'], course_id, title, '', '', 'concluded', 'import',
             rec['结论'], ts, ts, ts))
        report_id = cur.lastrowid
        mcur = db.execute(
            'INSERT INTO measures (report_id, seq, content, verify_indicator, assignee_id, '
            'deadline, status) VALUES (?, 1, ?, ?, ?, ?, ?)',
            (report_id, rec['措施内容'], rec['验证指标'], teacher['id'],
             config.DEFAULT_MEASURE_DEADLINE, 'finished'))
        measure_id = mcur.lastrowid
        # 达成报告按措施一对一：历史报告视为已由责任人（教师）提交并审批通过，
        # 整体结论存于 reports.conclusion
        db.execute(
            'INSERT INTO achievements (measure_id, report_id, submitter_id, content, '
            "status, submitted_at) VALUES (?, ?, ?, ?, 'approved', ?)",
            (measure_id, report_id, teacher['id'], rec['措施内容'], ts))
        seen_in_batch.add((rec['教师工号'], rec['课程代码'], rec['学年'], rec['学期']))
        result['success_rows'] += 1

    return result
