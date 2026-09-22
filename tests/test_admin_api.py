# -*- coding: utf-8 -*-
"""管理端 API 级测试：Flask test client + 临时数据库（用户/学年/课程/审批归属）"""
import io
import os
import sqlite3
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config  # noqa: E402
from werkzeug.security import generate_password_hash  # noqa: E402

_TMPDIR = tempfile.TemporaryDirectory()
config.DATABASE = os.path.join(_TMPDIR.name, 'test.db')
config.UPLOAD_DIR = os.path.join(_TMPDIR.name, 'uploads')
# 默认关闭登录限速，避免共享的内存态限速器跨用例累计触发 429；
# 专门的限速用例会在其 setUp 中临时开启并 reset。
config.RATELIMIT_ENABLED = False

import db as db_mod  # noqa: E402
from app import create_app  # noqa: E402
from db import now  # noqa: E402
from ratelimit import limiter  # noqa: E402

PASSWORD = '123456'
# 显式密码（≥8 位）：用于需要"非随机初始密码、无强制改密标记"的既有用例
EXPLICIT_PWD = 'Explicit123'


def make_docx(paragraphs):
    """内存构造最小 docx（与 test_backend 中相同）"""
    buf = io.BytesIO()
    body = ''.join(
        f'<w:p><w:r><w:t xml:space="preserve">{p}</w:t></w:r></w:p>'
        for p in paragraphs)
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w='
           '"http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:body>{body}</w:body></w:document>')
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types/>')
        z.writestr('word/document.xml', xml)
    buf.seek(0)
    return buf


def _conn():
    conn = sqlite3.connect(config.DATABASE)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    return conn


class AdminApiBase(unittest.TestCase):
    def setUp(self):
        for suffix in ('', '-wal', '-shm'):
            p = config.DATABASE + suffix
            if os.path.exists(p):
                os.remove(p)
        db_mod.init_db()
        conn = _conn()
        try:
            users = [
                ('teacher01', '王老师', 1, 0, 0, 0),
                ('teacher02', '吴老师', 1, 0, 0, 0),
                ('leader01', '张主任', 0, 1, 0, 0),
                ('leader02', '刘主任', 0, 1, 0, 0),
                ('admin01', '管理员', 0, 0, 1, 0),
                ('dual01', '双角色', 1, 1, 0, 0),
                ('disabled01', '停用教师', 1, 0, 0, 1),
            ]
            for u, n, t, l, a, d in users:
                conn.execute(
                    'INSERT INTO users (username, password_hash, real_name, is_teacher, '
                    'is_leader, is_admin, is_disabled) VALUES (?, ?, ?, ?, ?, ?, ?)',
                    (u, generate_password_hash(PASSWORD), n, t, l, a, d))
            conn.execute("INSERT INTO academic_years (name) VALUES ('2025-2026学年')")
            leader = conn.execute(
                "SELECT id FROM users WHERE username='leader01'").fetchone()
            conn.execute(
                'INSERT INTO courses (name, code, academic_year, term, start_date, '
                'end_date, leader_id) VALUES (?, ?, ?, ?, ?, ?, ?)',
                ('工程制图', 'GCTZ1001', '2025-2026学年', '第一学期',
                 '2025-09-01', '2025-12-31', leader['id']))
            conn.commit()
        finally:
            conn.close()
        self.app = create_app()
        self.app.config['TESTING'] = True
        self.client = self.app.test_client()

    # ---------- 工具 ----------

    def login(self, username, password=PASSWORD):
        return self.client.post('/api/auth/login',
                                json={'username': username, 'password': password})

    def uid(self, username):
        conn = _conn()
        try:
            return conn.execute(
                'SELECT id FROM users WHERE username=?', (username,)).fetchone()['id']
        finally:
            conn.close()

    def query(self, sql, args=()):
        conn = _conn()
        try:
            return conn.execute(sql, args).fetchall()
        finally:
            conn.close()

    def exec_sql(self, sql, args=()):
        conn = _conn()
        try:
            cur = conn.execute(sql, args)
            conn.commit()
            return cur.lastrowid
        finally:
            conn.close()

    def insert_report(self, teacher, leader_username, status):
        """直接写库构造指定负责人的待审报告，返回报告 id"""
        teacher_id = self.uid(teacher)
        leader_id = self.uid(leader_username) if leader_username else None
        course_id = self.query('SELECT id FROM courses')[0]['id']
        return self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, status, '
            'source, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
            (teacher_id, course_id, leader_id, f'报告-{status}', status,
             'upload', now(), now()))


class AuthFlagTest(AdminApiBase):
    def test_disabled_login_rejected(self):
        r = self.login('disabled01')
        self.assertEqual(r.status_code, 403)
        self.assertIn('停用', r.get_json()['error'])

    def test_disabled_session_rejected(self):
        # 管理员新建用户 → 新用户登录保持会话 → 管理员停用 → 旧会话访问被拒并清除
        self.login('admin01')
        r = self.client.post('/api/admin/users', json={
            'username': 'tmp01', 'real_name': '临时', 'is_teacher': 1,
            'password': EXPLICIT_PWD})
        self.assertTrue(r.get_json()['ok'])
        uid = r.get_json()['data']['id']

        user_client = self.app.test_client()
        self.assertEqual(
            user_client.post('/api/auth/login',
                             json={'username': 'tmp01', 'password': EXPLICIT_PWD}).status_code, 200)
        self.assertEqual(user_client.get('/api/reports').status_code, 200)

        self.client.patch(f'/api/admin/users/{uid}', json={'is_disabled': 1})
        r = user_client.get('/api/reports')
        self.assertEqual(r.status_code, 403)
        self.assertIn('停用', r.get_json()['error'])
        # 会话已被清除，再次访问回到未登录（且停用账号无法重新登录）
        self.assertEqual(user_client.get('/api/reports').status_code, 401)
        self.assertEqual(self.login('tmp01', EXPLICIT_PWD).status_code, 403)

    def test_login_returns_flags_not_role(self):
        data = self.login('dual01').get_json()['data']
        self.assertNotIn('role', data)
        self.assertTrue(data['is_teacher'])
        self.assertTrue(data['is_leader'])
        self.assertFalse(data['is_admin'])


class UserManageTest(AdminApiBase):
    def test_non_admin_cannot_access(self):
        self.login('teacher01')
        self.assertEqual(self.client.get('/api/admin/users').status_code, 403)

    def test_create_user_random_password_and_force_change(self):
        self.login('admin01')
        r = self.client.post('/api/admin/users', json={
            'username': 'new01', 'real_name': '新教师', 'is_teacher': 1})
        self.assertTrue(r.get_json()['ok'], r.get_json())
        data = r.get_json()['data']
        # 未显式指定密码 → 返回随机初始密码，且标记需首次登录修改
        initial = data.get('initial_password')
        self.assertTrue(initial)
        self.assertNotEqual(initial, '123456')
        self.assertGreaterEqual(len(initial), 8)
        self.assertTrue(data['must_change_password'])

        # 用随机初始密码可登录（登录本身不被 428 拦截）
        user_client = self.app.test_client()
        lr = user_client.post('/api/auth/login',
                              json={'username': 'new01', 'password': initial})
        self.assertEqual(lr.status_code, 200)
        self.assertTrue(lr.get_json()['data']['must_change_password'])

        # 改密前访问业务接口被 428 拦截
        self.assertEqual(user_client.get('/api/reports').status_code, 428)

        # 改密（旧=初始，新≥8）成功后解除强制改密
        cr = user_client.put('/api/user/password',
                             json={'old_password': initial, 'new_password': 'NewPass123'})
        self.assertTrue(cr.get_json()['ok'], cr.get_json())
        self.assertEqual(user_client.get('/api/reports').status_code, 200)

        # 过短的新密码被拒（最短 8 位）
        short = user_client.put('/api/user/password',
                                json={'old_password': 'NewPass123', 'new_password': '1234567'})
        self.assertEqual(short.status_code, 400)
        self.assertIn('至少8位', short.get_json()['error'])

    def test_create_user_explicit_short_password_rejected(self):
        self.login('admin01')
        r = self.client.post('/api/admin/users', json={
            'username': 'short01', 'real_name': '短密码', 'is_teacher': 1,
            'password': '1234567'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('至少8位', r.get_json()['error'])

    def test_create_user_duplicate_and_no_role(self):
        self.login('admin01')
        r = self.client.post('/api/admin/users', json={
            'username': 'teacher01', 'real_name': '重复', 'is_teacher': 1})
        self.assertEqual(r.status_code, 400)
        self.assertIn('已存在', r.get_json()['error'])
        r = self.client.post('/api/admin/users', json={
            'username': 'norole01', 'real_name': '无角色'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('至少分配一个角色', r.get_json()['error'])

    def test_update_user_and_self_protection(self):
        self.login('admin01')
        tid = self.uid('teacher01')
        r = self.client.patch(f'/api/admin/users/{tid}',
                              json={'real_name': '王老师改', 'is_leader': 1})
        data = r.get_json()['data']
        self.assertEqual(data['real_name'], '王老师改')
        self.assertTrue(data['is_leader'])
        # 不能改自己
        aid = self.uid('admin01')
        r = self.client.patch(f'/api/admin/users/{aid}', json={'real_name': 'X'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('不能修改自己的账号信息', r.get_json()['error'])

    def test_last_admin_protection(self):
        self.login('admin01')
        a2 = self.client.post('/api/admin/users', json={
            'username': 'admin02', 'real_name': '管理员2', 'is_admin': 1,
            'password': EXPLICIT_PWD}).get_json()['data']['id']
        # 两个启用管理员时取消其一合法（系统仍有 admin01）
        r = self.client.patch(f'/api/admin/users/{a2}', json={'is_admin': 0})
        self.assertTrue(r.get_json()['ok'], r.get_json())
        self.client.patch(f'/api/admin/users/{a2}', json={'is_admin': 1})

        # admin02 停用 admin01 → 系统仅剩 admin02 一个启用管理员；
        # admin01 会话立即失效，admin02 自我降级被自我保护拦截（纵深防御）
        client2 = self.app.test_client()
        client2.post('/api/auth/login',
                     json={'username': 'admin02', 'password': EXPLICIT_PWD})
        r = client2.patch(f"/api/admin/users/{self.uid('admin01')}", json={'is_disabled': 1})
        self.assertTrue(r.get_json()['ok'])
        r = self.client.get('/api/admin/users')
        self.assertEqual(r.status_code, 403)
        r = client2.patch(f'/api/admin/users/{a2}', json={'is_admin': 0})
        self.assertEqual(r.status_code, 400)
        self.assertIn('不能修改自己的账号信息', r.get_json()['error'])
        # 对已停用的 admin01 取消管理权限不触发保护（其已非启用中管理员）
        r = client2.patch(f"/api/admin/users/{self.uid('admin01')}",
                          json={'is_admin': 0})
        self.assertTrue(r.get_json()['ok'])

    def test_reset_password(self):
        self.exec_sql('UPDATE users SET password_hash=? WHERE username=?',
                      (generate_password_hash('oldpass'), 'teacher01'))
        self.assertEqual(self.login('teacher01', 'oldpass').status_code, 200)
        self.login('admin01')
        r = self.client.post(f"/api/admin/users/{self.uid('teacher01')}/reset-password")
        self.assertTrue(r.get_json()['ok'])
        # 重置为随机新密码：原密码失效，返回的 new_password 可登录
        newpwd = r.get_json()['data']['new_password']
        self.assertTrue(newpwd)
        self.assertNotEqual(newpwd, '123456')
        self.assertGreaterEqual(len(newpwd), 8)
        self.assertEqual(self.login('teacher01', 'oldpass').status_code, 401)
        self.assertEqual(self.login('teacher01', newpwd).status_code, 200)
        # 重置后标记需首次登录修改
        row = self.query('SELECT must_change_password FROM users WHERE username=?',
                         ('teacher01',))[0]
        self.assertEqual(row['must_change_password'], 1)

    def test_delete_user_double_confirm_and_cascade(self):
        tid = self.uid('teacher01')
        t2 = self.uid('teacher02')
        # 构造 teacher01 的报告链：报告 + 措施 + 达成
        os.makedirs(config.UPLOAD_DIR, exist_ok=True)
        fpath = os.path.join(config.UPLOAD_DIR, 'tobedeleted.docx')
        with open(fpath, 'wb') as f:
            f.write(b'x')
        rid = self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, file_path, '
            'status, source, created_at, updated_at) '
            'VALUES (?, 1, NULL, ?, ?, ?, ?, ?, ?)',
            (tid, '待删报告', fpath, 'executing', 'upload', now(), now()))
        mid = self.exec_sql(
            'INSERT INTO measures (report_id, seq, content, status) '
            'VALUES (?, 1, ?, ?)', (rid, '措施A', 'assigned'))
        self.exec_sql(
            'INSERT INTO achievements (measure_id, report_id, submitter_id, content, '
            'status, submitted_at) VALUES (?, ?, ?, ?, ?, ?)',
            (mid, rid, tid, '达成', 'submitted', now()))
        # teacher02 报告中措施指派给 teacher01；teacher02 报告的负责人快照为 teacher01
        rid2 = self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, status, '
            'source, created_at, updated_at) VALUES (?, 1, ?, ?, ?, ?, ?, ?)',
            (t2, tid, '他人报告', 'submitted', 'upload', now(), now()))
        mid2 = self.exec_sql(
            'INSERT INTO measures (report_id, seq, content, assignee_id, status) '
            'VALUES (?, 1, ?, ?, ?)', (rid2, '措施B', tid, 'assigned'))
        self.exec_sql(
            "INSERT INTO rejections (target_type, target_id, reason, rejecter_id, "
            f"created_at) VALUES ('measures', ?, ?, ?, ?)",
            (rid, '原因', tid, now()))
        self.exec_sql(
            "INSERT INTO action_log (user_id, action, object_type, object_id, detail, "
            f"created_at) VALUES (?, ?, 'report', ?, '{{}}', ?)",
            (tid, '旧操作', rid, now()))

        self.login('admin01')
        body_bad_name = {'confirm_username': 'wrong', 'admin_password': PASSWORD}
        r = self.client.delete(f'/api/admin/users/{tid}', json=body_bad_name)
        self.assertEqual(r.status_code, 400)
        self.assertIn('确认用户名不匹配', r.get_json()['error'])
        body_bad_pwd = {'confirm_username': 'teacher01', 'admin_password': 'wrong'}
        r = self.client.delete(f'/api/admin/users/{tid}', json=body_bad_pwd)
        self.assertEqual(r.status_code, 403)
        self.assertIn('管理员密码验证失败', r.get_json()['error'])
        # 禁止删自己
        r = self.client.delete(
            f"/api/admin/users/{self.uid('admin01')}",
            json={'confirm_username': 'admin01', 'admin_password': PASSWORD})
        self.assertEqual(r.status_code, 400)

        r = self.client.delete(f'/api/admin/users/{tid}', json={
            'confirm_username': 'teacher01', 'admin_password': PASSWORD})
        self.assertTrue(r.get_json()['ok'], r.get_json())
        # 级联校验
        self.assertEqual(len(self.query('SELECT * FROM users WHERE id=?', (tid,))), 0)
        self.assertEqual(len(self.query('SELECT * FROM reports WHERE id=?', (rid,))), 0)
        self.assertEqual(len(self.query('SELECT * FROM measures WHERE report_id=?', (rid,))), 0)
        self.assertEqual(len(self.query('SELECT * FROM achievements WHERE report_id=?', (rid,))), 0)
        self.assertEqual(
            len(self.query('SELECT * FROM rejections WHERE rejecter_id=?', (tid,))), 0)
        self.assertIsNone(self.query(
            'SELECT assignee_id FROM measures WHERE id=?', (mid2,))[0]['assignee_id'])
        self.assertIsNone(self.query(
            'SELECT leader_id FROM reports WHERE id=?', (rid2,))[0]['leader_id'])
        self.assertIsNone(self.query(
            'SELECT user_id FROM action_log WHERE action=?', ('旧操作',))[0]['user_id'])
        self.assertFalse(os.path.exists(fpath))

    def test_delete_user_blocked_when_assignee_of_executing(self):
        # 被删教师是执行中报告的措施责任人 → 阻断删除（非 500），不删任何数据
        tid = self.uid('teacher01')
        t2 = self.uid('teacher02')
        rid = self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, status, '
            'source, created_at, updated_at) VALUES (?, 1, NULL, ?, ?, ?, ?, ?)',
            (t2, '执行中报告', 'executing', 'upload', now(), now()))
        self.exec_sql(
            'INSERT INTO measures (report_id, seq, content, assignee_id, deadline, status) '
            'VALUES (?, 1, ?, ?, ?, ?)', (rid, '措施A', tid, '2099-12-31', 'assigned'))
        self.login('admin01')
        r = self.client.delete(f'/api/admin/users/{tid}', json={
            'confirm_username': 'teacher01', 'admin_password': PASSWORD})
        self.assertEqual(r.status_code, 400)
        self.assertIn('执行中报告的措施责任人', r.get_json()['error'])
        # 未删除任何数据
        self.assertEqual(len(self.query('SELECT * FROM users WHERE id=?', (tid,))), 1)
        self.assertEqual(
            len(self.query('SELECT * FROM measures WHERE report_id=?', (rid,))), 1)

    def test_delete_user_cleans_submitted_achievement_no_fk(self):
        # 被删教师曾在他人已定级报告上提交达成（submitter_id NOT NULL FK）：
        # 非执行中牵连→不阻断，删除前先清理其达成的 submitter 引用，不报 FK 500
        tid = self.uid('teacher01')
        t2 = self.uid('teacher02')
        rid = self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, status, '
            'source, created_at, updated_at) VALUES (?, 1, NULL, ?, ?, ?, ?, ?)',
            (t2, '已定级报告', 'concluded', 'upload', now(), now()))
        mid = self.exec_sql(
            'INSERT INTO measures (report_id, seq, content, assignee_id, status) '
            'VALUES (?, 1, ?, ?, ?)', (rid, '措施A', tid, 'finished'))
        self.exec_sql(
            'INSERT INTO achievements (measure_id, report_id, submitter_id, content, '
            'status, submitted_at) VALUES (?, ?, ?, ?, ?, ?)',
            (mid, rid, tid, '历史达成', 'approved', now()))
        self.login('admin01')
        r = self.client.delete(f'/api/admin/users/{tid}', json={
            'confirm_username': 'teacher01', 'admin_password': PASSWORD})
        self.assertTrue(r.get_json()['ok'], r.get_json())
        # 用户已删、其提交的达成被清理（FK 满足），他人报告本身保留
        self.assertEqual(len(self.query('SELECT * FROM users WHERE id=?', (tid,))), 0)
        self.assertEqual(
            len(self.query('SELECT * FROM achievements WHERE submitter_id=?', (tid,))), 0)
        self.assertEqual(len(self.query('SELECT * FROM reports WHERE id=?', (rid,))), 1)

    def test_list_users_contains_flags(self):
        self.login('admin01')
        rows = self.client.get('/api/admin/users').get_json()['data']
        self.assertEqual(len(rows), 7)
        for row in rows:
            for f in ('is_teacher', 'is_leader', 'is_admin', 'is_disabled'):
                self.assertIn(f, row)


class YearManageTest(AdminApiBase):
    def test_years_list_with_course_count(self):
        self.login('admin01')
        rows = self.client.get('/api/admin/years').get_json()['data']
        self.assertEqual(rows[0]['name'], '2025-2026学年')
        self.assertEqual(rows[0]['course_count'], 1)

    def test_delete_referenced_year_rejected(self):
        self.login('admin01')
        yid = self.query('SELECT id FROM academic_years')[0]['id']
        r = self.client.delete(f'/api/admin/years/{yid}')
        self.assertEqual(r.status_code, 400)
        self.assertIn('课程', r.get_json()['error'])

    def test_create_duplicate_and_delete_unused_year(self):
        self.login('admin01')
        r = self.client.post('/api/admin/years', json={'name': '2025-2026学年'})
        self.assertEqual(r.status_code, 400)
        r = self.client.post('/api/admin/years', json={'name': '2026-2027学年'})
        self.assertTrue(r.get_json()['ok'])
        yid = r.get_json()['data']['id']
        r = self.client.delete(f'/api/admin/years/{yid}')
        self.assertTrue(r.get_json()['ok'])


class CourseManageTest(AdminApiBase):
    def _post(self, body):
        return self.client.post('/api/admin/courses', json=body)

    def test_create_validations(self):
        self.login('admin01')
        base = {'name': '新课程', 'code': 'NEW1001', 'academic_year': '2025-2026学年',
                'term': '第一学期', 'start_date': '2025-09-01', 'end_date': '2025-12-31'}
        # 非法学期
        r = self._post({**base, 'term': '第三学期'})
        self.assertEqual(r.status_code, 400)
        # 非法日期格式（端点接受 YYYY-MM-DD 与 YYYY/M/D，点分隔为非法）
        r = self._post({**base, 'start_date': '2025.09.01'})
        self.assertEqual(r.status_code, 400)
        # 开始晚于结束
        r = self._post({**base, 'start_date': '2026-01-01', 'end_date': '2025-12-31'})
        self.assertEqual(r.status_code, 400)
        # 学年不存在
        r = self._post({**base, 'academic_year': '1999-2000学年'})
        self.assertEqual(r.status_code, 400)
        # 唯一键重复
        r = self._post({**base, 'code': 'GCTZ1001'})
        self.assertEqual(r.status_code, 400)
        # 负责人不是负责人角色
        r = self._post({**base, 'leader_id': self.uid('teacher01')})
        self.assertEqual(r.status_code, 400)
        # 负责人已停用
        r = self._post({**base, 'leader_id': self.uid('disabled01')})
        self.assertEqual(r.status_code, 400)
        # 合法创建（含负责人）
        r = self._post({**base, 'leader_id': self.uid('leader02')})
        data = r.get_json()
        self.assertTrue(data['ok'], data)
        self.assertEqual(data['data']['leader_name'], '刘主任')

    def test_list_courses_and_patch_leader(self):
        self.login('admin01')
        rows = self.client.get('/api/admin/courses').get_json()['data']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['leader_name'], '张主任')
        self.assertEqual(rows[0]['report_count'], 0)

        cid = rows[0]['id']
        r = self.client.patch(f'/api/admin/courses/{cid}/leader',
                              json={'leader_id': self.uid('leader02')})
        self.assertEqual(r.get_json()['data']['leader_id'], self.uid('leader02'))
        # 允许 null 清空
        r = self.client.patch(f'/api/admin/courses/{cid}/leader', json={'leader_id': None})
        self.assertIsNone(r.get_json()['data']['leader_id'])
        # 非法负责人被拒
        r = self.client.patch(f'/api/admin/courses/{cid}/leader',
                              json={'leader_id': self.uid('teacher02')})
        self.assertEqual(r.status_code, 400)

    def _build_course_with_report(self):
        """在唯一课程下构造一份带措施/达成/退回记录及上传文件的报告，返回 (cid, rid, path)"""
        cid = self.query('SELECT id FROM courses')[0]['id']
        os.makedirs(config.UPLOAD_DIR, exist_ok=True)
        path = os.path.join(config.UPLOAD_DIR, 'cascade_test.docx')
        with open(path, 'wb') as f:
            f.write(b'x')
        rid = self.exec_sql(
            'INSERT INTO reports (teacher_id, course_id, leader_id, title, status, '
            'source, file_path, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (self.uid('teacher01'), cid, self.uid('leader01'), '级联报告', 'approved',
             'upload', path, now(), now()))
        mid = self.exec_sql(
            'INSERT INTO measures (report_id, seq, content, status) VALUES (?, 1, ?, ?)',
            (rid, '措施内容', 'assigned'))
        aid = self.exec_sql(
            'INSERT INTO achievements (measure_id, report_id, submitter_id, content, '
            'status, submitted_at) VALUES (?, ?, ?, ?, ?, ?)',
            (mid, rid, self.uid('teacher01'), '达成', 'approved', now()))
        self.exec_sql(
            "INSERT INTO rejections (target_type, target_id, reason, rejecter_id, "
            "created_at) VALUES ('measures', ?, '退回', ?, ?)",
            (rid, self.uid('leader01'), now()))
        self.exec_sql(
            "INSERT INTO rejections (target_type, target_id, reason, rejecter_id, "
            "created_at) VALUES ('achievement', ?, '退回', ?, ?)",
            (aid, self.uid('leader01'), now()))
        return cid, rid, path

    def test_delete_course_requires_name_code_confirmation(self):
        cid, _rid, _path = self._build_course_with_report()
        self.login('admin01')
        # 名称正确但编号错误 → 拒绝
        r = self.client.delete(f'/api/admin/courses/{cid}',
                               json={'name': '工程制图', 'code': 'WRONG'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('确认', r.get_json()['error'])
        # 课程仍存在
        self.assertEqual(len(self.query('SELECT id FROM courses WHERE id=?', (cid,))), 1)

    def test_delete_course_cascades_reports_and_files(self):
        cid, rid, path = self._build_course_with_report()
        self.login('admin01')
        r = self.client.delete(f'/api/admin/courses/{cid}',
                               json={'name': '工程制图', 'code': 'GCTZ1001'})
        data = r.get_json()
        self.assertTrue(data['ok'], data)
        self.assertEqual(data['data']['deleted_reports'], 1)
        # 课程与报告及关联行全部删除
        self.assertEqual(self.query('SELECT id FROM courses WHERE id=?', (cid,)), [])
        self.assertEqual(self.query('SELECT id FROM reports WHERE id=?', (rid,)), [])
        self.assertEqual(self.query('SELECT id FROM measures WHERE report_id=?', (rid,)), [])
        self.assertEqual(self.query('SELECT id FROM achievements WHERE report_id=?', (rid,)), [])
        self.assertEqual(
            self.query("SELECT id FROM rejections WHERE target_id=? ", (rid,)), [])
        # 上传文件已从磁盘删除
        self.assertFalse(os.path.exists(path))


class ApprovalScopeTest(AdminApiBase):
    def test_pending_filtered_by_leader(self):
        r1 = self.insert_report('teacher01', 'leader01', 'submitted')
        r2 = self.insert_report('teacher01', 'leader02', 'submitted')
        r3 = self.insert_report('teacher01', None, 'submitted')

        self.login('leader01')
        data = self.client.get('/api/leader/pending').get_json()['data']
        ids = {r['id'] for r in data['pending_measures']}
        self.assertIn(r1, ids)
        self.assertIn(r3, ids)      # 未指派：任何负责人可见
        self.assertNotIn(r2, ids)   # 他人指派：不可见

    def test_reject_scope_check(self):
        r1 = self.insert_report('teacher01', 'leader01', 'submitted')
        r3 = self.insert_report('teacher01', None, 'submitted')

        self.login('leader02')
        r = self.client.post(f'/api/reports/{r1}/reject', json={'reason': '不行'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('已指派给其他负责人审批', r.get_json()['error'])
        # 未指派报告任何负责人可审
        r = self.client.post(f'/api/reports/{r3}/reject', json={'reason': '不行'})
        self.assertTrue(r.get_json()['ok'])

    def test_approve_scope_check(self):
        r1 = self.insert_report('teacher01', 'leader01', 'submitted')
        self.exec_sql('INSERT INTO measures (report_id, seq, content, status) '
                      'VALUES (?, 1, ?, ?)', (r1, '措施', 'draft'))
        self.login('leader02')
        r = self.client.post(f'/api/reports/{r1}/approve', json={
            'assignments': [{'measure_id': 1, 'assignee_id': self.uid('teacher01')}],
            'deadlines': {'1': '2026-06-30'}})
        self.assertEqual(r.status_code, 403)


class DualRoleTest(AdminApiBase):
    def test_dual_role_accesses_both_sides(self):
        self.login('dual01')
        # 负责人端点可访问
        r = self.client.get('/api/leader/pending')
        self.assertEqual(r.status_code, 200)
        # 教师端点角色检查通过（缺参数 400，而非 403）
        r = self.client.post('/api/reports', data={})
        self.assertEqual(r.status_code, 400)

    def test_pure_leader_cannot_use_teacher_endpoint(self):
        self.login('leader01')
        r = self.client.post('/api/reports', data={})
        self.assertEqual(r.status_code, 403)

    def test_create_report_snapshots_course_leader(self):
        self.login('teacher01')
        course_id = self.query('SELECT id FROM courses')[0]['id']
        docx = make_docx(['课程目标达成情况', '1. 加强模型教学。', '次年验证指标：达成度不低于0.8。'])
        r = self.client.post('/api/reports', data={
            'course_id': str(course_id),
            'file': (docx, '报告.docx')}, content_type='multipart/form-data')
        data = r.get_json()
        self.assertTrue(data['ok'], data)
        row = self.query('SELECT leader_id FROM reports WHERE id=?',
                         (data['data']['id'],))[0]
        self.assertEqual(row['leader_id'], self.uid('leader01'))

    def test_dictionaries_flag_based(self):
        self.login('teacher01')
        data = self.client.get('/api/dictionaries').get_json()['data']
        teacher_names = {t['username'] for t in data['teachers']}
        self.assertIn('teacher01', teacher_names)
        self.assertNotIn('disabled01', teacher_names)   # 停用教师不进字典
        self.assertEqual(data['academic_years'], ['2025-2026学年'])
        for u in data['users']:
            self.assertNotIn('role', u)
            self.assertIn('is_teacher', u)
        course = data['courses'][0]
        self.assertEqual(course['leader_name'], '张主任')


class UserImportCredentialTest(AdminApiBase):
    CSV = ('用户名,姓名,角色,密码（留空自动生成随机密码）\n'
           'imp01,导入甲,教师,\n'
           'imp02,导入乙,负责人,\n'
           'imp03,导入丙,教师,Explicit123\n')

    def _import(self, csv_text, dry_run=False):
        data = {'file': (io.BytesIO(csv_text.encode('utf-8')), 'users.csv')}
        if dry_run:
            data['dry_run'] = '1'
        return self.client.post('/api/admin/users/import', data=data,
                                content_type='multipart/form-data')

    def test_import_generates_credential_bundle(self):
        self.login('admin01')
        r = self._import(self.CSV)
        body = r.get_json()
        self.assertTrue(body['ok'], body)
        self.assertEqual(body['data']['success_count'], 3)
        bundle = body['data'].get('credentials_bundle')
        self.assertTrue(bundle)

        # 下载凭据表并校验内容
        dl = self.client.get(f'/api/admin/users/import/credentials/{bundle}')
        self.assertEqual(dl.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(dl.data), read_only=True, data_only=True)
        ws = wb.active
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        wb.close()
        self.assertEqual(set(rows.keys()), {'imp01', 'imp02', 'imp03'})
        # 显式密码原样保留；留空者生成随机密码（≥8 位，非 123456）
        self.assertEqual(rows['imp03'][2], 'Explicit123')
        for uname in ('imp01', 'imp02'):
            pwd = rows[uname][2]
            self.assertTrue(pwd)
            self.assertGreaterEqual(len(pwd), 8)
            self.assertNotEqual(pwd, '123456')

        # 留空密码的用户被标记需首次登录修改；显式密码者不标记
        self.assertEqual(self.query(
            'SELECT must_change_password FROM users WHERE username=?', ('imp01',))[0][0], 1)
        self.assertEqual(self.query(
            'SELECT must_change_password FROM users WHERE username=?', ('imp03',))[0][0], 0)

    def test_import_dry_run_no_bundle(self):
        self.login('admin01')
        r = self._import(self.CSV, dry_run=True)
        body = r.get_json()
        self.assertTrue(body['ok'], body)
        self.assertEqual(body['data']['success_count'], 3)
        self.assertNotIn('credentials_bundle', body['data'])
        self.assertIn('preview_rows', body['data'])
        # 预览未真正建号
        self.assertEqual(len(self.query(
            'SELECT id FROM users WHERE username=?', ('imp01',))), 0)

    def test_import_short_explicit_password_fails_row(self):
        self.login('admin01')
        csv_text = ('用户名,姓名,角色,密码\n'
                    'ok01,甲,教师,\n'
                    'bad01,乙,教师,123\n')
        r = self._import(csv_text)
        body = r.get_json()
        self.assertTrue(body['ok'], body)
        self.assertEqual(body['data']['success_count'], 1)
        self.assertEqual(body['data']['failed_count'], 1)
        self.assertIn('至少8位', body['data']['failed_rows'][0]['error'])

    def test_credential_download_requires_admin(self):
        self.login('admin01')
        bundle = self._import(self.CSV).get_json()['data']['credentials_bundle']
        # 非管理员不可下载
        self.login('teacher01')
        self.assertEqual(
            self.client.get(f'/api/admin/users/import/credentials/{bundle}').status_code, 403)

    def test_credential_download_unknown_bundle_404(self):
        self.login('admin01')
        self.assertEqual(
            self.client.get('/api/admin/users/import/credentials/deadbeef').status_code, 404)


class SecurityConfigTest(AdminApiBase):
    def test_get_defaults(self):
        self.login('admin01')
        data = self.client.get('/api/admin/security-config').get_json()['data']
        self.assertTrue(data['force_pwd_change'])
        self.assertEqual(data['login_max_attempts'], 10)
        self.assertEqual(data['login_lock_minutes'], 15)
        self.assertEqual(data['min_password_len'], 8)

    def test_put_persists(self):
        self.login('admin01')
        r = self.client.put('/api/admin/security-config', json={
            'force_pwd_change': False, 'login_max_attempts': 5, 'login_lock_minutes': 30})
        self.assertTrue(r.get_json()['ok'], r.get_json())
        data = r.get_json()['data']
        self.assertFalse(data['force_pwd_change'])
        self.assertEqual(data['login_max_attempts'], 5)
        self.assertEqual(data['login_lock_minutes'], 30)
        # 再次读取确认已落库
        again = self.client.get('/api/admin/security-config').get_json()['data']
        self.assertFalse(again['force_pwd_change'])
        self.assertEqual(again['login_max_attempts'], 5)

    def test_put_clamps_out_of_range(self):
        self.login('admin01')
        r = self.client.put('/api/admin/security-config', json={
            'login_max_attempts': 9999, 'login_lock_minutes': 0})
        data = r.get_json()['data']
        self.assertEqual(data['login_max_attempts'], 100)   # 夹取上限
        self.assertEqual(data['login_lock_minutes'], 1)      # 夹取下限

    def test_non_admin_forbidden(self):
        self.login('teacher01')
        self.assertEqual(self.client.get('/api/admin/security-config').status_code, 403)

    def test_force_change_off_allows_access_without_change(self):
        # 关闭强制改密后，持随机初始密码的用户登录即可直接访问业务接口（不 428）
        self.login('admin01')
        self.client.put('/api/admin/security-config', json={'force_pwd_change': False})
        initial = self.client.post('/api/admin/users', json={
            'username': 'nc01', 'real_name': '免改密', 'is_teacher': 1}
        ).get_json()['data']['initial_password']
        uc = self.app.test_client()
        self.assertEqual(uc.post('/api/auth/login',
                                 json={'username': 'nc01', 'password': initial}).status_code, 200)
        self.assertEqual(uc.get('/api/reports').status_code, 200)


class LoginRateLimitTest(AdminApiBase):
    """专门验证登录限速：临时开启 RATELIMIT_ENABLED 并 reset 共享限速器。"""

    def setUp(self):
        config.RATELIMIT_ENABLED = True
        super().setUp()
        limiter.reset()
        # 收紧阈值便于测试：3 次 / 15 分钟
        self.exec_sql(
            "INSERT INTO settings (key, value) VALUES ('security_login_max_attempts', '3') "
            "ON CONFLICT(key) DO UPDATE SET value='3'")
        self.exec_sql(
            "INSERT INTO settings (key, value) VALUES ('security_login_lock_minutes', '15') "
            "ON CONFLICT(key) DO UPDATE SET value='15'")

    def tearDown(self):
        limiter.reset()
        config.RATELIMIT_ENABLED = False

    def test_login_blocked_after_max_attempts(self):
        codes = [self.client.post('/api/auth/login', json={
            'username': 'teacher01', 'password': 'wrong'}).status_code for _ in range(5)]
        # 前 3 次进入视图（401），第 4、5 次被限速（429）
        self.assertEqual(codes[:3], [401, 401, 401])
        self.assertEqual(codes[3:], [429, 429])
        # 锁定期内即使密码正确也被拒
        locked = self.client.post('/api/auth/login',
                                  json={'username': 'teacher01', 'password': PASSWORD})
        self.assertEqual(locked.status_code, 429)

    def test_limit_is_per_username(self):
        # 耗尽 teacher01 的配额
        for _ in range(4):
            self.client.post('/api/auth/login',
                             json={'username': 'teacher01', 'password': 'wrong'})
        # 另一用户名不受影响（独立计数键）
        other = self.client.post('/api/auth/login',
                                 json={'username': 'teacher02', 'password': PASSWORD})
        self.assertEqual(other.status_code, 200)


if __name__ == '__main__':
    unittest.main()