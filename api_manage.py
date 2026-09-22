# -*- coding: utf-8 -*-
"""管理端蓝图：用户/学年/课程管理、AI 配置（全部要求 is_admin，写操作均记审计日志）"""
import csv
import io
import json
import os
import urllib.error
import urllib.request
import uuid
from datetime import datetime, timedelta

from flask import Blueprint, current_app, request
from werkzeug.security import check_password_hash, generate_password_hash

import config
from auth import ok, fail, role_required
from db import get_db, log_action, now
from services import ai_config
from services import dingtalk_config
from services import security_config

bp = Blueprint('manage', __name__, url_prefix='/api/admin')

ROLE_FLAGS = ('is_teacher', 'is_leader', 'is_admin')
TERMS = ('第一学期', '第二学期')

# 导入凭据包有效期（分钟）：过期后不可再下载
CREDENTIAL_BUNDLE_TTL_MIN = 30

def _term_names(count):
    """根据学期数量生成学期名称列表"""
    cn = ['零', '一', '二', '三', '四', '五', '六', '七', '八', '九', '十']
    return [f'第{cn[i]}学期' for i in range(1, count + 1)]
# 测试连接超时上限（秒），与实际解析超时取小
AI_TEST_MAX_TIMEOUT = 10.0


def _flag(v):
    return 1 if v else 0


def _user_dict(r):
    keys = r.keys()
    dt_bound = bool(r['dingtalk_userid']) if 'dingtalk_userid' in keys else False
    must_change = bool(r['must_change_password']) if 'must_change_password' in keys else False
    return {'id': r['id'], 'username': r['username'], 'real_name': r['real_name'],
            'is_teacher': bool(r['is_teacher']), 'is_leader': bool(r['is_leader']),
            'is_admin': bool(r['is_admin']), 'is_disabled': bool(r['is_disabled']),
            'dingtalk_bound': dt_bound, 'must_change_password': must_change}


def _parse_date(s):
    if isinstance(s, datetime):
        return s.date()
    if not isinstance(s, str):
        return None
    s = s.strip()
    if not s:
        return None
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ---------- 用户管理 ----------

@bp.get('/users')
@role_required('admin')
def list_users():
    db = get_db()
    rows = db.execute('SELECT * FROM users ORDER BY id').fetchall()
    return ok([_user_dict(r) for r in rows])


@bp.post('/users')
@role_required('admin')
def create_user():
    user = request.current_user
    body = request.get_json(silent=True) or {}
    username = str(body.get('username') or '').strip()
    real_name = str(body.get('real_name') or '').strip()
    if not username:
        return fail('用户名不能为空')
    if not real_name:
        return fail('姓名不能为空')
    flags = {f: _flag(body.get(f)) for f in ROLE_FLAGS}
    if not any(flags.values()):
        return fail('请至少分配一个角色')

    # 未显式提供密码则生成随机初始密码并强制首次登录修改
    provided = str(body.get('password') or '').strip()
    if provided:
        if len(provided) < config.MIN_PASSWORD_LEN:
            return fail(f'密码长度至少{config.MIN_PASSWORD_LEN}位')
        password = provided
        must_change = 0
    else:
        password = security_config.random_password()
        must_change = 1

    db = get_db()
    exists = db.execute('SELECT id FROM users WHERE username=?', (username,)).fetchone()
    if exists:
        return fail('用户名已存在')
    cur = db.execute(
        'INSERT INTO users (username, password_hash, real_name, is_teacher, is_leader, '
        'is_admin, must_change_password) VALUES (?, ?, ?, ?, ?, ?, ?)',
        (username, generate_password_hash(password), real_name,
         flags['is_teacher'], flags['is_leader'], flags['is_admin'], must_change))
    new_id = cur.lastrowid
    log_action(db, user['id'], '新增用户', 'user', new_id,
               {'username': username, 'real_name': real_name,
                'is_teacher': flags['is_teacher'], 'is_leader': flags['is_leader'],
                'is_admin': flags['is_admin'], 'random_password': bool(must_change)})
    db.commit()
    row = db.execute('SELECT * FROM users WHERE id=?', (new_id,)).fetchone()
    data = _user_dict(row)
    if must_change:
        data['initial_password'] = password
    return ok(data)


# ---------- 角色解析辅助 ----------

_ROLE_MAP = {
    '教师': 'is_teacher', '教师端': 'is_teacher',
    '负责人': 'is_leader', '专业负责人': 'is_leader', '负责人端': 'is_leader',
    '管理员': 'is_admin',
}

def _parse_roles(role_str):
    """解析角色字符串，返回 {is_teacher: 0/1, is_leader: 0/1, is_admin: 0/1}"""
    flags = {'is_teacher': 0, 'is_leader': 0, 'is_admin': 0}
    if not role_str:
        return flags
    parts = [p.strip() for p in role_str.replace(',', '，').split('，') if p.strip()]
    for p in parts:
        key = _ROLE_MAP.get(p)
        if key:
            flags[key] = 1
    return flags


@bp.post('/users/import')
@role_required('admin')
def import_users():
    user = request.current_user
    file = request.files.get('file')
    if file is None or not file.filename:
        return fail('请上传文件')
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ('.xlsx', '.xlsm', '.csv'):
        return fail('仅支持 .xlsx / .csv 文件')

    # 解析文件
    rows = []
    try:
        if ext == '.csv':
            text = file.stream.read().decode('utf-8-sig')
            reader = csv.reader(io.StringIO(text))
            header = None
            for r in reader:
                if header is None:
                    header = [c.strip() for c in r]
                    continue
                if any(c.strip() for c in r):
                    rows.append(dict(zip(header, [c.strip() for c in r])))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            ws = wb.active
            header = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = [str(c).strip() if c is not None else '' for c in row]
                if i == 0:
                    header = cells
                    continue
                if any(cells):
                    rows.append(dict(zip(header, cells)))
            wb.close()
    except Exception as e:
        return fail(f'文件解析失败：{e}')

    if not rows:
        return fail('文件中无有效数据行')

    # 列名映射：支持中英文表头
    COL_MAP = {
        '用户名': 'username', 'username': 'username', '登录名': 'username', '工号': 'username',
        '姓名': 'real_name', 'real_name': 'real_name', '真实姓名': 'real_name',
        '角色': 'role', 'role': 'role', '工作端': 'role', '角色（教师，负责人，管理员）': 'role',
        '密码': 'password', 'password': 'password', '初始密码': 'password',
        '密码（留空自动生成随机密码）': 'password', '密码（留空默认123456）': 'password',
    }

    db = get_db()
    success_count = 0
    failed_rows = []
    preview_rows = []
    credentials = []  # [{username, real_name, password, random}]
    dry_run = request.form.get('dry_run', '') in ('1', 'true', 'yes')

    for idx, row in enumerate(rows, start=2):  # 从第2行开始（第1行是表头）
        # 规范化列名
        normalized = {}
        for k, v in row.items():
            mapped = COL_MAP.get(k.strip() if k else '')
            if mapped:
                normalized[mapped] = v

        username = normalized.get('username', '').strip()
        real_name = normalized.get('real_name', '').strip()
        role_str = normalized.get('role', '').strip()
        provided = normalized.get('password', '').strip()

        # 校验
        if not username:
            failed_rows.append({'row': idx, 'error': '用户名为空'})
            continue
        if not real_name:
            failed_rows.append({'row': idx, 'error': '姓名为空'})
            continue
        flags = _parse_roles(role_str)
        if not any(flags.values()):
            failed_rows.append({'row': idx, 'error': f'角色无效：{role_str}'})
            continue
        if provided and len(provided) < config.MIN_PASSWORD_LEN:
            failed_rows.append({'row': idx,
                                'error': f'密码长度至少{config.MIN_PASSWORD_LEN}位'})
            continue

        # 检查重复
        exists = db.execute('SELECT id FROM users WHERE username=?', (username,)).fetchone()
        if exists:
            failed_rows.append({'row': idx, 'error': f'用户名已存在：{username}'})
            continue

        # 未填密码则生成随机初始密码并强制首次登录修改
        if provided:
            password, must_change = provided, 0
        else:
            password, must_change = security_config.random_password(), 1

        if dry_run:
            roles = []
            if flags['is_teacher']: roles.append('教师')
            if flags['is_leader']: roles.append('负责人')
            if flags['is_admin']: roles.append('管理员')
            preview_rows.append({'row': idx, 'username': username,
                                'real_name': real_name, 'role': '、'.join(roles),
                                'random_password': bool(must_change)})
            success_count += 1
        else:
            try:
                db.execute(
                    'INSERT INTO users (username, password_hash, real_name, is_teacher, '
                    'is_leader, is_admin, must_change_password) VALUES (?, ?, ?, ?, ?, ?, ?)',
                    (username, generate_password_hash(password), real_name,
                     flags['is_teacher'], flags['is_leader'], flags['is_admin'],
                     must_change))
                credentials.append({'username': username, 'real_name': real_name,
                                    'password': password, 'random': bool(must_change)})
                success_count += 1
            except Exception as e:
                failed_rows.append({'row': idx, 'error': str(e)})

    bundle_id = None
    if not dry_run:
        if credentials:
            bundle_id = uuid.uuid4().hex
            expires = (datetime.now() + timedelta(minutes=CREDENTIAL_BUNDLE_TTL_MIN)
                       ).strftime('%Y-%m-%d %H:%M:%S')
            db.execute(
                'INSERT INTO credential_bundles (id, data, expires_at) VALUES (?, ?, ?)',
                (bundle_id, json.dumps(credentials, ensure_ascii=False), expires))
        log_action(db, user['id'], '批量导入用户', 'user', 0,
                   {'success_count': success_count, 'failed_count': len(failed_rows),
                    'filename': file.filename, 'bundle': bundle_id})
        db.commit()

    result = {
        'success_count': success_count,
        'failed_count': len(failed_rows),
        'failed_rows': failed_rows[:20]
    }
    if dry_run:
        result['preview_rows'] = preview_rows
    elif bundle_id:
        result['credentials_bundle'] = bundle_id
    return ok(result)


@bp.get('/users/import/template')
@role_required('admin')
def users_import_template():
    """下载用户批量导入模板"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '用户导入'
    ws.append(['用户名', '姓名', '角色（教师，负责人，管理员）', '密码（留空自动生成随机密码）'])
    ws.append(['teacher03', '李老师', '教师', ''])
    ws.append(['leader03', '张主任', '负责人', ''])
    ws.append(['teacher04', '王老师', '教师，负责人', ''])
    # 设置列宽
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()

    from flask import send_file
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='用户导入模板.xlsx')


@bp.get('/users/import/credentials/<bundle_id>')
@role_required('admin')
def download_import_credentials(bundle_id):
    """下载某次批量导入生成的用户名/初始密码对照表（限时、仅管理员）。"""
    user = request.current_user
    db = get_db()
    row = db.execute('SELECT data, expires_at FROM credential_bundles WHERE id=?',
                     (bundle_id,)).fetchone()
    if row is None:
        return fail('凭据不存在或已过期', 404)
    if row['expires_at'] < now():
        db.execute('DELETE FROM credential_bundles WHERE id=?', (bundle_id,))
        db.commit()
        return fail('凭据已过期，请重新导入后下载', 410)

    try:
        credentials = json.loads(row['data'])
    except (ValueError, TypeError):
        return fail('凭据数据损坏', 500)

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '初始密码'
    ws.append(['用户名', '姓名', '初始密码'])
    for c in credentials:
        ws.append([c.get('username', ''), c.get('real_name', ''),
                   c.get('password', '')])
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()

    log_action(db, user['id'], '下载导入凭据', 'user', 0, {'bundle': bundle_id})
    db.commit()

    from flask import send_file
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name='用户初始密码.xlsx')

@bp.patch('/users/<int:uid>')
@role_required('admin')
def update_user(uid):
    user = request.current_user
    if uid == user['id']:
        return fail('不能修改自己的账号信息')
    db = get_db()
    target = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
    if target is None:
        return fail('用户不存在', 404)

    body = request.get_json(silent=True) or {}
    fields, values = [], []
    if 'real_name' in body:
        real_name = str(body['real_name'] or '').strip()
        if not real_name:
            return fail('姓名不能为空')
        fields.append('real_name=?')
        values.append(real_name)
    for f in ROLE_FLAGS + ('is_disabled',):
        if f in body:
            fields.append(f'{f}=?')
            values.append(_flag(body[f]))
    if not fields:
        return fail('没有可修改的字段')

    # 保护：不允许让系统中最后一个启用中的管理员失去可用管理员身份
    new_is_admin = _flag(body['is_admin']) if 'is_admin' in body else target['is_admin']
    new_is_disabled = _flag(body['is_disabled']) if 'is_disabled' in body else target['is_disabled']
    if target['is_admin'] and not target['is_disabled'] \
            and not (new_is_admin and not new_is_disabled):
        other = db.execute(
            'SELECT COUNT(*) AS c FROM users '
            'WHERE is_admin=1 AND is_disabled=0 AND id!=?', (uid,)).fetchone()['c']
        if other == 0:
            return fail('不能取消系统中最后一个启用管理员的管理权限')

    db.execute(f"UPDATE users SET {', '.join(fields)} WHERE id=?", values + [uid])
    log_action(db, user['id'], '修改用户', 'user', uid,
               dict(zip((f[:-2] for f in fields), values)))
    db.commit()
    row = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
    return ok(_user_dict(row))


@bp.post('/users/<int:uid>/reset-password')
@role_required('admin')
def reset_password(uid):
    user = request.current_user
    db = get_db()
    target = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
    if target is None:
        return fail('用户不存在', 404)
    new_password = security_config.random_password()
    db.execute('UPDATE users SET password_hash=?, must_change_password=1 WHERE id=?',
               (generate_password_hash(new_password), uid))
    log_action(db, user['id'], '重置用户密码', 'user', uid,
               {'username': target['username'], 'random_password': True})
    db.commit()
    return ok({'id': uid, 'username': target['username'],
               'new_password': new_password})


@bp.post('/users/<int:uid>/unbind-dingtalk')
@role_required('admin')
def admin_unbind_dingtalk(uid):
    """管理员解绑用户的钉钉账号"""
    user = request.current_user
    db = get_db()
    target = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
    if target is None:
        return fail('用户不存在', 404)
    keys = target.keys()
    if 'dingtalk_userid' not in keys or not target['dingtalk_userid']:
        return fail('该用户未绑定钉钉')
    db.execute("UPDATE users SET dingtalk_userid='' WHERE id=?", (uid,))
    log_action(db, user['id'], '管理员解绑用户钉钉', 'user', uid,
               {'username': target['username']})
    db.commit()
    return ok()


@bp.delete('/users/<int:uid>')
@role_required('admin')
def delete_user(uid):
    user = request.current_user
    if uid == user['id']:
        return fail('不能删除自己的账号')
    db = get_db()
    target = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
    if target is None:
        return fail('用户不存在', 404)

    body = request.get_json(silent=True) or {}
    confirm_username = str(body.get('confirm_username') or '')
    if confirm_username != target['username']:
        return fail('确认用户名不匹配')
    admin_password = str(body.get('admin_password') or '')
    if not check_password_hash(user['password_hash'], admin_password):
        return fail('管理员密码验证失败', 403)

    # 阻断保护：若该教师是任何"执行中"报告措施的责任人(assignee_id)，删除会使这些报告
    # 的措施悬空、无法推进（孤儿死锁），故要求先由负责人完成或改派后再删账号，不删任何数据。
    executing_reports = db.execute(
        'SELECT COUNT(DISTINCT r.id) AS c FROM reports r '
        'JOIN measures m ON m.report_id=r.id '
        "WHERE r.status='executing' AND m.assignee_id=?",
        (uid,)).fetchone()['c']
    if executing_reports:
        return fail(f'该教师是 {executing_reports} 份执行中报告的措施责任人，'
                    '请先由专业负责人完成或改派这些报告后再删除账号')

    # 单事务级联清理：其报告的达成/措施/报告 → 其退回记录 → 置空外部引用 → 删用户
    reports = db.execute('SELECT id, file_path FROM reports WHERE teacher_id=?',
                         (uid,)).fetchall()
    report_ids = [r['id'] for r in reports]
    # 先清理其在任意报告（含他人/历史/已定级报告）上提交的达成，满足
    # achievements.submitter_id NOT NULL REFERENCES users(id) 外键，避免删除用户时 FK 失败 500
    db.execute('DELETE FROM achievements WHERE submitter_id=?', (uid,))
    if report_ids:
        ph = ','.join('?' * len(report_ids))
        db.execute(f'DELETE FROM achievements WHERE report_id IN ({ph})', report_ids)
        db.execute(f'DELETE FROM measures WHERE report_id IN ({ph})', report_ids)
        db.execute(f'DELETE FROM reports WHERE id IN ({ph})', report_ids)
    db.execute('DELETE FROM rejections WHERE rejecter_id=?', (uid,))
    db.execute('UPDATE measures SET assignee_id=NULL WHERE assignee_id=?', (uid,))
    db.execute('UPDATE courses SET leader_id=NULL WHERE leader_id=?', (uid,))
    db.execute('UPDATE reports SET leader_id=NULL WHERE leader_id=?', (uid,))
    db.execute('UPDATE action_log SET user_id=NULL WHERE user_id=?', (uid,))
    db.execute('DELETE FROM users WHERE id=?', (uid,))
    log_action(db, user['id'], '删除用户', 'user', uid,
               {'username': target['username'], 'deleted_reports': len(report_ids)})
    db.commit()

    # 磁盘上传文件删除：失败仅记录，不阻断
    for r in reports:
        path = r['file_path'] or ''
        if path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError as e:
                current_app.logger.warning('删除上传文件失败 %s: %s', path, e)
    return ok({'id': uid, 'username': target['username']})


# ---------- 学年管理 ----------

@bp.get('/years')
@role_required('admin')
def list_years():
    db = get_db()
    rows = db.execute(
        'SELECT ay.id, ay.name, ay.semester_count, COUNT(c.id) AS course_count '
        'FROM academic_years ay LEFT JOIN courses c ON c.academic_year=ay.name '
        'GROUP BY ay.id, ay.name ORDER BY ay.name').fetchall()
    return ok([{'id': r['id'], 'name': r['name'], 'course_count': r['course_count'],
                'semester_count': r['semester_count']}
               for r in rows])


@bp.post('/years')
@role_required('admin')
def create_year():
    user = request.current_user
    body = request.get_json(silent=True) or {}
    name = str(body.get('name') or '').strip()
    semester_count = body.get('semester_count', 2)
    try:
        semester_count = int(semester_count)
    except (ValueError, TypeError):
        semester_count = 2
    if semester_count < 1 or semester_count > 10:
        return fail('学期数量必须在1-10之间')
    if not name:
        return fail('学年名称不能为空')
    db = get_db()
    exists = db.execute('SELECT id FROM academic_years WHERE name=?', (name,)).fetchone()
    if exists:
        return fail('学年已存在')
    cur = db.execute('INSERT INTO academic_years (name, semester_count) VALUES (?, ?)',
                     (name, semester_count))
    new_id = cur.lastrowid
    log_action(db, user['id'], '新增学年', 'academic_year', new_id,
               {'name': name, 'semester_count': semester_count})
    db.commit()
    return ok({'id': new_id, 'name': name, 'semester_count': semester_count})


@bp.post('/years/import')
@role_required('admin')
def import_years():
    user = request.current_user
    file = request.files.get('file')
    if file is None or not file.filename:
        return fail('请上传文件')
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ('.xlsx', '.xlsm', '.csv'):
        return fail('仅支持 .xlsx / .csv 文件')

    # 解析文件
    rows = []
    try:
        if ext == '.csv':
            text = file.stream.read().decode('utf-8-sig')
            reader = csv.reader(io.StringIO(text))
            header = None
            for r in reader:
                if header is None:
                    header = [c.strip() for c in r]
                    continue
                if any(c.strip() for c in r):
                    rows.append(dict(zip(header, [c.strip() for c in r])))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            ws = wb.active
            header = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = [str(c).strip() if c is not None else '' for c in row]
                if i == 0:
                    header = cells
                    continue
                if any(cells):
                    rows.append(dict(zip(header, cells)))
            wb.close()
    except Exception as e:
        return fail(f'文件解析失败：{e}')

    if not rows:
        return fail('文件中无有效数据行')

    # 列名映射
    COL_MAP = {
        '学年名称': 'name', '学年': 'name', 'name': 'name',
        '学期数': 'semester_count', '学期数量': 'semester_count', 'semester_count': 'semester_count',
        '学期数（留空默认2学期）': 'semester_count',
    }

    db = get_db()
    success_count = 0
    failed_rows = []
    preview_rows = []
    dry_run = request.form.get('dry_run', '') in ('1', 'true', 'yes')

    for idx, row in enumerate(rows, start=2):
        normalized = {}
        for k, v in row.items():
            mapped = COL_MAP.get(k.strip() if k else '')
            if mapped:
                normalized[mapped] = v

        name = normalized.get('name', '').strip()
        sc_str = normalized.get('semester_count', '').strip()
        try:
            semester_count = int(sc_str) if sc_str else 2
        except (ValueError, TypeError):
            semester_count = 2

        if not name:
            failed_rows.append({'row': idx, 'error': '学年名称为空'})
            continue
        if semester_count < 1 or semester_count > 10:
            failed_rows.append({'row': idx, 'error': f'学期数无效：{sc_str}（需1-10）'})
            continue

        exists = db.execute('SELECT id FROM academic_years WHERE name=?', (name,)).fetchone()
        if exists:
            failed_rows.append({'row': idx, 'error': f'学年已存在：{name}'})
            continue

        if dry_run:
            preview_rows.append({'row': idx, 'name': name,
                                'semester_count': semester_count})
            success_count += 1
        else:
            try:
                db.execute('INSERT INTO academic_years (name, semester_count) VALUES (?, ?)',
                           (name, semester_count))
                success_count += 1
            except Exception as e:
                failed_rows.append({'row': idx, 'error': str(e)})

    if not dry_run:
        log_action(db, user['id'], '批量导入学年', 'academic_year', 0,
                   {'success_count': success_count, 'failed_count': len(failed_rows),
                    'filename': file.filename})
        db.commit()

    result = {
        'success_count': success_count,
        'failed_count': len(failed_rows),
        'failed_rows': failed_rows[:20]
    }
    if dry_run:
        result['preview_rows'] = preview_rows
    return ok(result)


@bp.get('/years/import/template')
@role_required('admin')
def years_import_template():
    """下载学年批量导入模板"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '学年导入'
    ws.append(['学年名称', '学期数（留空默认2学期）'])
    ws.append(['2025-2026学年', '2'])
    ws.append(['2026-2027学年', ''])
    ws.append(['2027-2028学年', '3'])
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()

    from flask import send_file
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='学年导入模板.xlsx')


@bp.delete('/years/<int:yid>')
@role_required('admin')
def delete_year(yid):
    user = request.current_user
    db = get_db()
    year = db.execute('SELECT * FROM academic_years WHERE id=?', (yid,)).fetchone()
    if year is None:
        return fail('学年不存在', 404)
    used = db.execute(
        'SELECT COUNT(*) AS c FROM courses WHERE academic_year=?',
        (year['name'],)).fetchone()['c']
    if used:
        return fail(f'该学年下有 {used} 门课程，不能删除')
    db.execute('DELETE FROM academic_years WHERE id=?', (yid,))
    log_action(db, user['id'], '删除学年', 'academic_year', yid, {'name': year['name']})
    db.commit()
    return ok({'id': yid, 'name': year['name']})


# ---------- 课程管理 ----------

@bp.get('/courses')
@role_required('admin')
def list_courses():
    db = get_db()
    rows = db.execute(
        'SELECT c.*, u.real_name AS leader_name, '
        '(SELECT COUNT(*) FROM reports r WHERE r.course_id=c.id) AS report_count '
        'FROM courses c LEFT JOIN users u ON c.leader_id=u.id '
        'ORDER BY c.academic_year, c.term, c.code').fetchall()
    items = []
    for r in rows:
        item = dict(r)
        for k in ('start_date', 'end_date'):
            if item[k] and len(item[k]) > 10:
                item[k] = item[k][:10]
        items.append(item)
    return ok(items)


def _validate_leader(db, leader_id):
    """校验负责人：启用中且 is_leader=1，返回 (错误信息, 行)"""
    leader = db.execute('SELECT * FROM users WHERE id=?', (leader_id,)).fetchone()
    if leader is None:
        return '负责人不存在', None
    if not leader['is_leader']:
        return '负责人必须具有负责人角色', None
    if leader['is_disabled']:
        return '负责人账号已停用', None
    return None, leader


@bp.post('/courses')
@role_required('admin')
def create_course():
    user = request.current_user
    body = request.get_json(silent=True) or {}
    name = str(body.get('name') or '').strip()
    code = str(body.get('code') or '').strip()
    academic_year = str(body.get('academic_year') or '').strip()
    term = str(body.get('term') or '').strip()
    start_date = str(body.get('start_date') or '').strip()
    end_date = str(body.get('end_date') or '').strip()
    leader_id = body.get('leader_id')

    if not name or not code or not academic_year or not term:
        return fail('课程名称、代码、学年、学期均不能为空')

    db = get_db()
    year_row = db.execute('SELECT id, semester_count FROM academic_years WHERE name=?',
                          (academic_year,)).fetchone()
    if year_row is None:
        return fail('学年不存在')
    valid_terms = _term_names(year_row['semester_count'])
    if term not in valid_terms:
        return fail(f"学期无效，可选：{'/'.join(valid_terms)}")

    sd = _parse_date(start_date) if start_date else None
    ed = _parse_date(end_date) if end_date else None
    if (start_date and sd is None) or (end_date and ed is None):
        return fail('起止日期格式应为 YYYY-MM-DD 或 YYYY/M/D')
    if sd and ed and sd > ed:
        return fail('开始日期不能晚于结束日期')

    dup = db.execute(
        'SELECT id FROM courses WHERE code=? AND academic_year=? AND term=?',
        (code, academic_year, term)).fetchone()
    if dup:
        return fail('同代码同学年同学期的课程已存在')

    if leader_id is not None:
        try:
            leader_id = int(leader_id)
        except (TypeError, ValueError):
            return fail('负责人编号非法')
        err, _ = _validate_leader(db, leader_id)
        if err:
            return fail(err)

    cur = db.execute(
        'INSERT INTO courses (name, code, academic_year, term, start_date, end_date, leader_id) '
        'VALUES (?, ?, ?, ?, ?, ?, ?)',
        (name, code, academic_year, term,
         str(sd) if sd else None, str(ed) if ed else None, leader_id))
    new_id = cur.lastrowid
    log_action(db, user['id'], '新增课程', 'course', new_id,
               {'name': name, 'code': code, 'academic_year': academic_year,
                'term': term, 'leader_id': leader_id})
    db.commit()
    row = db.execute(
        'SELECT c.*, u.real_name AS leader_name FROM courses c '
        'LEFT JOIN users u ON c.leader_id=u.id WHERE c.id=?', (new_id,)).fetchone()
    return ok(dict(row))


@bp.patch('/courses/<int:cid>/leader')
@role_required('admin')
def update_course_leader(cid):
    user = request.current_user
    db = get_db()
    course = db.execute('SELECT * FROM courses WHERE id=?', (cid,)).fetchone()
    if course is None:
        return fail('课程不存在', 404)

    body = request.get_json(silent=True) or {}
    leader_id = body.get('leader_id')
    if leader_id is not None:
        try:
            leader_id = int(leader_id)
        except (TypeError, ValueError):
            return fail('负责人编号非法')
        err, _ = _validate_leader(db, leader_id)
        if err:
            return fail(err)

    db.execute('UPDATE courses SET leader_id=? WHERE id=?', (leader_id, cid))
    log_action(db, user['id'], '设置课程负责人', 'course', cid,
               {'course': f"{course['name']}({course['code']})", 'leader_id': leader_id})
    db.commit()
    row = db.execute(
        'SELECT c.*, u.real_name AS leader_name FROM courses c '
        'LEFT JOIN users u ON c.leader_id=u.id WHERE c.id=?', (cid,)).fetchone()
    return ok(dict(row))


@bp.delete('/courses/<int:cid>')
@role_required('admin')
def delete_course(cid):
    """删除课程，级联删除其下所有报告（不限状态）及报告的措施/达成/退回记录与上传文件。

    请求体须回传课程 name 与 code 二次确认，防止误删；单事务级联删除，
    磁盘上传文件删除失败仅记录不阻断；action_log 保留（审计不回删）。
    """
    user = request.current_user
    body = request.get_json(silent=True) or {}
    confirm_name = str(body.get('name') or '').strip()
    confirm_code = str(body.get('code') or '').strip()

    db = get_db()
    course = db.execute('SELECT * FROM courses WHERE id=?', (cid,)).fetchone()
    if course is None:
        return fail('课程不存在', 404)
    if confirm_name != (course['name'] or '').strip() or \
            confirm_code != (course['code'] or '').strip():
        return fail('请输入正确的课程名称与编号以确认删除')

    reports = db.execute(
        'SELECT id, file_path FROM reports WHERE course_id=?', (cid,)).fetchall()
    report_ids = [r['id'] for r in reports]
    if report_ids:
        ph = ','.join('?' * len(report_ids))
        db.execute(
            "DELETE FROM rejections WHERE target_type='achievement' AND target_id IN "
            f'(SELECT id FROM achievements WHERE report_id IN ({ph}))', report_ids)
        db.execute(f'DELETE FROM achievements WHERE report_id IN ({ph})', report_ids)
        db.execute(
            "DELETE FROM rejections WHERE target_type='measures' AND target_id IN "
            f'({ph})', report_ids)
        db.execute(f'DELETE FROM measures WHERE report_id IN ({ph})', report_ids)
        db.execute(f'DELETE FROM reports WHERE id IN ({ph})', report_ids)

    db.execute('DELETE FROM courses WHERE id=?', (cid,))
    log_action(db, user['id'], '删除课程', 'course', cid,
               {'name': course['name'], 'code': course['code'],
                'academic_year': course['academic_year'], 'term': course['term'],
                'deleted_reports': len(report_ids)})
    db.commit()

    # 磁盘上传文件删除：失败仅记录，不阻断
    for r in reports:
        path = r['file_path'] or ''
        if path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError as e:
                current_app.logger.warning('删除上传文件失败 %s: %s', path, e)
    return ok({'id': cid, 'name': course['name'], 'deleted_reports': len(report_ids)})


@bp.post('/courses/import')
@role_required('admin')
def import_courses():
    user = request.current_user
    file = request.files.get('file')
    if file is None or not file.filename:
        return fail('请上传文件')
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ('.xlsx', '.xlsm', '.csv'):
        return fail('仅支持 .xlsx / .csv 文件')

    rows = []
    try:
        if ext == '.csv':
            text = file.stream.read().decode('utf-8-sig')
            reader = csv.reader(io.StringIO(text))
            header = None
            for r in reader:
                if header is None:
                    header = [c.strip() for c in r]
                    continue
                if any(c.strip() for c in r):
                    rows.append(dict(zip(header, [c.strip() for c in r])))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            ws = wb.active
            header = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = [str(c).strip() if c is not None else '' for c in row]
                if i == 0:
                    header = cells
                    continue
                if any(cells):
                    rows.append(dict(zip(header, cells)))
            wb.close()
    except Exception as e:
        return fail(f'文件解析失败：{e}')

    if not rows:
        return fail('文件中无有效数据行')

    COL_MAP = {
        '课程名称': 'name', '名称': 'name', 'name': 'name',
        '课程编号': 'code', '编号': 'code', 'code': 'code',
        '学年': 'academic_year', 'academic_year': 'academic_year',
        '学期（填数字，如1表示第一学期）': 'term', '学期': 'term', 'term': 'term',
        '开始日期': 'start_date', 'start_date': 'start_date',
        '结束日期': 'end_date', 'end_date': 'end_date',
        '专业负责人': 'leader_username', '负责人': 'leader_username',
        'leader_username': 'leader_username',
    }

    db = get_db()
    success_count = 0
    failed_rows = []
    preview_rows = []
    dry_run = request.form.get('dry_run', '') in ('1', 'true', 'yes')

    for idx, row in enumerate(rows, start=2):
        normalized = {}
        for k, v in row.items():
            mapped = COL_MAP.get(k.strip() if k else '')
            if mapped:
                normalized[mapped] = v

        name = normalized.get('name', '').strip()
        code = normalized.get('code', '').strip()
        academic_year = normalized.get('academic_year', '').strip()
        term_str = normalized.get('term', '').strip()
        start_date = normalized.get('start_date', '').strip()
        end_date = normalized.get('end_date', '').strip()
        leader_username = normalized.get('leader_username', '').strip()

        if not name or not code or not academic_year or not term_str:
            failed_rows.append({'row': idx, 'error': '课程名称、编号、学年、学期不能为空'})
            continue

        try:
            term_num = int(term_str)
        except (ValueError, TypeError):
            failed_rows.append({'row': idx, 'error': f'学期应为数字，当前值：{term_str}'})
            continue

        year_row = db.execute('SELECT id, semester_count FROM academic_years WHERE name=?',
                              (academic_year,)).fetchone()
        if year_row is None:
            failed_rows.append({'row': idx, 'error': f'学年不存在：{academic_year}'})
            continue

        valid_terms = _term_names(year_row['semester_count'])
        if term_num < 1 or term_num > year_row['semester_count']:
            failed_rows.append({'row': idx, 'error': f'学期超出范围：{term_str}（该学年有{year_row["semester_count"]}个学期）'})
            continue
        term = valid_terms[term_num - 1]

        sd = _parse_date(start_date) if start_date else None
        ed = _parse_date(end_date) if end_date else None
        if (start_date and sd is None) or (end_date and ed is None):
            failed_rows.append({'row': idx, 'error': '起止日期格式应为 YYYY-MM-DD 或 YYYY/M/D'})
            continue
        if sd and ed and sd > ed:
            failed_rows.append({'row': idx, 'error': '开始日期不能晚于结束日期'})
            continue

        dup = db.execute(
            'SELECT id FROM courses WHERE code=? AND academic_year=? AND term=?',
            (code, academic_year, term)).fetchone()
        if dup:
            failed_rows.append({'row': idx, 'error': f'同代码同学年同学期的课程已存在：{code} / {academic_year} / {term}'})
            continue

        leader_id = None
        if leader_username:
            leader_row = db.execute(
                'SELECT id, is_leader, is_disabled FROM users WHERE username=?',
                (leader_username,)).fetchone()
            if leader_row is None:
                failed_rows.append({'row': idx, 'error': f'负责人用户名不存在：{leader_username}'})
                continue
            if not leader_row['is_leader']:
                failed_rows.append({'row': idx, 'error': f'用户 {leader_username} 不具有负责人角色'})
                continue
            if leader_row['is_disabled']:
                failed_rows.append({'row': idx, 'error': f'用户 {leader_username} 已停用'})
                continue
            leader_id = leader_row['id']

        if dry_run:
            preview_rows.append({'row': idx, 'name': name, 'code': code,
                                'academic_year': academic_year, 'term': term,
                                'start_date': str(sd) if sd else (start_date or '—'),
                                'end_date': str(ed) if ed else (end_date or '—'),
                                'leader': leader_username or '—'})
            success_count += 1
        else:
            try:
                db.execute(
                    'INSERT INTO courses (name, code, academic_year, term, start_date, end_date, leader_id) '
                    'VALUES (?, ?, ?, ?, ?, ?, ?)',
                    (name, code, academic_year, term,
                     str(sd) if sd else None, str(ed) if ed else None, leader_id))
                success_count += 1
            except Exception as e:
                failed_rows.append({'row': idx, 'error': str(e)})

    if not dry_run:
        log_action(db, user['id'], '批量导入课程', 'course', 0,
                   {'success_count': success_count, 'failed_count': len(failed_rows),
                    'filename': file.filename})
        db.commit()

    result = {
        'success_count': success_count,
        'failed_count': len(failed_rows),
        'failed_rows': failed_rows[:20]
    }
    if dry_run:
        result['preview_rows'] = preview_rows
    return ok(result)


@bp.get('/courses/import/template')
@role_required('admin')
def courses_import_template():
    """下载课程批量导入模板"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '课程导入'
    ws.append(['课程名称', '课程编号', '学年', '学期（填数字，如1表示第一学期）', '开始日期', '结束日期', '专业负责人'])
    ws.append(['工程制图', 'GCTZ1001', '2025-2026学年', '1', '2025-09-01', '2025-12-31', 'leader01'])
    ws.append(['高等数学', 'GDSX3001', '2025-2026学年', '2', '2026-02-28', '2026-06-30', ''])
    for col_letter, width in [('A', 16), ('B', 14), ('C', 20), ('D', 35), ('E', 14), ('F', 14), ('G', 16)]:
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()

    from flask import send_file
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='课程导入模板.xlsx')


# ---------- AI 配置 ----------

def _upsert_setting(db, key, value):
    db.execute(
        'INSERT INTO settings (key, value) VALUES (?, ?) '
        'ON CONFLICT(key) DO UPDATE SET value=excluded.value', (key, value))


def _ai_config_data(settings):
    """GET/PUT 统一返回结构：API Key 只写不读，仅返回是否已设置与掩码，绝不回显明文"""
    key_set = bool((settings['api_key'] or '').strip())
    return {
        'enabled': settings['enabled'],
        'api_key_set': key_set,
        'api_key_masked': ai_config.mask_api_key(settings['api_key']),
        'base_url': settings['base_url'],
        'model': settings['model'],
        'timeout': settings['timeout'],
        'prompt': settings['prompt'],
        'default_prompt': ai_config.DEFAULT_PROMPT,
        'effective_source': settings['source'],
    }


@bp.get('/ai-config')
@role_required('admin')
def get_ai_config():
    settings = ai_config.resolve_from_db(get_db())
    return ok(_ai_config_data(settings))


@bp.put('/ai-config')
@role_required('admin')
def put_ai_config():
    user = request.current_user
    db = get_db()
    body = request.get_json(silent=True) or {}

    changes = {}
    if 'enabled' in body:
        changes['ai_enabled'] = '1' if body.get('enabled') else '0'
    api_key_provided = 'api_key' in body and str(body.get('api_key') or '').strip()
    if api_key_provided:
        # api_key 非空才覆盖（写入后不可读回）
        changes['ai_api_key'] = str(body['api_key']).strip()
    for field, key in (('base_url', 'ai_base_url'), ('model', 'ai_model'),
                       ('prompt', 'ai_prompt')):
        if field in body:
            changes[key] = str(body.get(field) or '').strip()
    if 'timeout' in body:
        try:
            timeout = float(body.get('timeout'))
        except (TypeError, ValueError):
            return fail('超时时间必须为正数')
        if timeout <= 0:
            return fail('超时时间必须为正数')
        changes['ai_timeout'] = str(timeout)
    clear_key = bool(body.get('clear_api_key'))

    if not changes and not clear_key:
        return fail('没有可保存的配置项')

    for key, value in changes.items():
        _upsert_setting(db, key, value)
    if clear_key:
        # 清空 DB 密钥（回落环境变量）
        _upsert_setting(db, 'ai_api_key', '')

    # 审计日志不含密钥明文，仅记录是否设置了密钥与各字段来源
    detail = {'fields': sorted(changes.keys()), 'api_key_set': bool(api_key_provided),
              'clear_api_key': clear_key}
    log_action(db, user['id'], '修改 AI 配置', 'settings', None, detail)
    db.commit()
    return ok(_ai_config_data(ai_config.resolve_from_db(db)))


@bp.post('/ai-config/test')
@role_required('admin')
def test_ai_config():
    """用解析后的配置向 {base_url}/chat/completions 发最小请求验证连通性；
    成功/失败响应与日志均不含密钥明文"""
    user = request.current_user
    db = get_db()
    settings = ai_config.resolve_from_db(db)
    if not settings['enabled']:
        return fail('AI 功能未启用，请先启用并保存配置')
    if not (settings['api_key'] or '').strip():
        return fail('未配置 API Key，无法测试连接')

    url = (settings['base_url'] or '').rstrip('/') + '/chat/completions'
    body = {'model': settings['model'],
            'messages': [{'role': 'user', 'content': '回复OK'}]}
    data = json.dumps(body, ensure_ascii=False).encode('utf-8')
    req = urllib.request.Request(url, data=data, method='POST')
    req.add_header('Authorization', 'Bearer ' + settings['api_key'])
    req.add_header('Content-Type', 'application/json')
    try:
        timeout = min(float(settings['timeout'] or 0) or AI_TEST_MAX_TIMEOUT,
                      AI_TEST_MAX_TIMEOUT)
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read().decode('utf-8', errors='replace')
    except urllib.error.HTTPError as e:
        # 原始错误体仅写服务端日志（可能含密钥片段/内部信息），不回传前端
        err_body = ''
        try:
            err_body = e.read().decode('utf-8', errors='replace').strip()
        except Exception:  # noqa: BLE001 - 读不到响应体不影响日志记录
            pass
        current_app.logger.warning(
            'AI 连接测试失败 HTTP %s: %s', e.code, err_body[:500])
        return fail(f'AI 服务返回错误（HTTP {e.code}），请检查 API Key 与模型配置')
    except TimeoutError:
        return fail('连接超时，请检查 Base URL 与网络')
    except urllib.error.URLError as e:
        reason = getattr(e, 'reason', e)
        return fail(f'无法连接 AI 服务：{reason}')
    except OSError as e:
        return fail(f'连接失败：{type(e).__name__}')

    try:
        payload = json.loads(raw)
        model = payload.get('model') or settings['model']
        content = payload['choices'][0]['message']['content']
        if not str(content or '').strip():
            raise ValueError('empty content')
    except (ValueError, TypeError, KeyError, IndexError):
        return fail('AI 服务响应格式异常，请检查 Base URL 与模型名称')

    log_action(db, user['id'], '测试 AI 连接成功', 'settings', None,
               {'model': model})
    db.commit()
    return ok({'message': '连接成功', 'model': model})


# ---------- 钉钉登录配置 ----------

def _dingtalk_config_data(settings):
    """GET/PUT 统一返回结构：app_secret 只写不读"""
    secret_set = bool((settings['app_secret'] or '').strip())
    return {
        'enabled': settings['enabled'],
        'app_key': settings['app_key'],
        'app_secret_set': secret_set,
        'redirect_uri': settings['redirect_uri'],
        'effective_source': settings['source'],
    }


@bp.get('/dingtalk-config')
@role_required('admin')
def get_dingtalk_config():
    settings = dingtalk_config.resolve_from_db(get_db())
    return ok(_dingtalk_config_data(settings))


@bp.put('/dingtalk-config')
@role_required('admin')
def put_dingtalk_config():
    user = request.current_user
    db = get_db()
    body = request.get_json(silent=True) or {}

    changes = {}
    for field, key in (('app_key', 'dingtalk_app_key'),
                       ('redirect_uri', 'dingtalk_redirect_uri')):
        if field in body:
            changes[key] = str(body.get(field) or '').strip()
    secret_provided = 'app_secret' in body and str(body.get('app_secret') or '').strip()
    if secret_provided:
        changes['dingtalk_app_secret'] = str(body['app_secret']).strip()
    clear_secret = bool(body.get('clear_app_secret'))
    if clear_secret:
        changes['dingtalk_app_secret'] = ''

    for key, value in changes.items():
        _upsert_setting(db, key, value)

    detail = {k: ('***' if 'secret' in k else v) for k, v in changes.items()}
    log_action(db, user['id'], '修改钉钉登录配置', 'settings', None, detail)
    db.commit()

    settings = dingtalk_config.resolve_from_db(db)
    return ok(_dingtalk_config_data(settings))


# ---------- 安全选项配置 ----------

@bp.get('/security-config')
@role_required('admin')
def get_security_config():
    cfg = security_config.resolve_from_db(get_db())
    return ok({
        'force_pwd_change': cfg['force_pwd_change'],
        'login_max_attempts': cfg['login_max_attempts'],
        'login_lock_minutes': cfg['login_lock_minutes'],
        'min_password_len': config.MIN_PASSWORD_LEN,
        'attempts_range': [security_config.ATTEMPTS_MIN, security_config.ATTEMPTS_MAX],
        'lock_range': [security_config.LOCK_MINUTES_MIN,
                       security_config.LOCK_MINUTES_MAX],
    })


@bp.put('/security-config')
@role_required('admin')
def put_security_config():
    user = request.current_user
    db = get_db()
    body = request.get_json(silent=True) or {}

    changes = {}
    if 'force_pwd_change' in body:
        changes['security_force_pwd_change'] = '1' if body.get('force_pwd_change') else '0'
    if 'login_max_attempts' in body:
        attempts = security_config._parse_int(
            body.get('login_max_attempts'), security_config.DEFAULT_LOGIN_MAX_ATTEMPTS,
            security_config.ATTEMPTS_MIN, security_config.ATTEMPTS_MAX)
        changes['security_login_max_attempts'] = str(attempts)
    if 'login_lock_minutes' in body:
        lock = security_config._parse_int(
            body.get('login_lock_minutes'), security_config.DEFAULT_LOGIN_LOCK_MINUTES,
            security_config.LOCK_MINUTES_MIN, security_config.LOCK_MINUTES_MAX)
        changes['security_login_lock_minutes'] = str(lock)

    if not changes:
        return fail('没有可保存的配置项')

    for key, value in changes.items():
        _upsert_setting(db, key, value)
    log_action(db, user['id'], '修改安全选项', 'settings', None,
               {'fields': sorted(changes.keys())})
    db.commit()

    cfg = security_config.resolve_from_db(db)
    return ok({
        'force_pwd_change': cfg['force_pwd_change'],
        'login_max_attempts': cfg['login_max_attempts'],
        'login_lock_minutes': cfg['login_lock_minutes'],
        'min_password_len': config.MIN_PASSWORD_LEN,
        'attempts_range': [security_config.ATTEMPTS_MIN, security_config.ATTEMPTS_MAX],
        'lock_range': [security_config.LOCK_MINUTES_MIN,
                       security_config.LOCK_MINUTES_MAX],
    })
